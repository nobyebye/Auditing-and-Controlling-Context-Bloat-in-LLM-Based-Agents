"""Build the formal thesis draft directly from the retained UH DOTX template."""

from __future__ import annotations

import argparse
import copy
import hashlib
import re
import shutil
from dataclasses import dataclass
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from instantiate_dotx import instantiate
from build_thesis_core_figures import build as build_core_figures


TITLE = (
    "Detecting, Measuring, and Mitigating Context Bloat in LLM-Based Agents: "
    "A Runtime Auditing Approach"
)
AUTHOR = "Guochen Li"
PROGRAMME = ""
THESIS_TYPE = "Master's thesis"
MONTH_YEAR = ""
PLACE = "Helsinki"
SUPERVISOR = ""
KEYWORDS = (
    "LLM agents, context bloat, runtime auditing, context provenance, "
    "retrieval-augmented generation, agent memory, tool use"
)
TEMPLATE_SHA256 = (
    "88E740A2E6EFE25102AF5E29812A34C7B512590D499263E612CF7509FB2D805D"
)

CHAPTER_FILES = (
    "01_introduction.md",
    "02_background_related_work.md",
    "03_model_and_framework.md",
    "04_methodology.md",
    "05_results.md",
    "06_mitigation_discussion.md",
    "07_conclusion.md",
)

TABLE_CAPTIONS = {
    "02_background_related_work.md": [
        "Research areas related to runtime auditing of context bloat.",
    ],
    "04_methodology.md": [
        "Controlled experimental conditions and their workflow-specific realizations.",
    ],
    "05_results.md": [
        "Controlled detection results by bloat label.",
        "Context metrics by experimental condition.",
        "Source-specific context bloat in the primary cohort.",
        "Primary results by agent implementation.",
    ],
    "06_mitigation_discussion.md": [
        "Mitigation outcomes by workflow family and agent implementation.",
    ],
}
TABLE_NUMBERS = {
    "02_background_related_work.md": ["2.1"],
    "04_methodology.md": ["4.1"],
    "05_results.md": ["5.1", "5.2", "5.3", "5.4"],
    "06_mitigation_discussion.md": ["6.1"],
}
CORE_FIGURES = {
    "research_overview": {
        "file": "figure_1_1_research_overview.png",
        "number": "1.1",
        "caption": (
            "Research overview linking automatically constructed context, "
            "context bloat, runtime auditing, source-aware control, and the "
            "four research questions."
        ),
        "alt": (
            "Retrieval, memory, tool outputs, and history feed model-visible "
            "context. Runtime auditing captures, attributes, detects, localizes, "
            "and measures bloat before source-aware mitigation is evaluated for "
            "efficiency and task utility."
        ),
    },
    "experimental_workflow": {
        "file": "figure_4_1_experimental_workflow.png",
        "number": "4.1",
        "caption": (
            "Study A/B/C workflow from shared request capture through separated "
            "injected, heuristic, human-reference, and counterfactual evidence."
        ),
        "alt": (
            "Study A tests controlled pipeline consistency, Study B uses "
            "independently annotated natural traces, and Study C compares "
            "counterfactual removal and budget-matched mitigation."
        ),
    },
    "graphical_abstract": {
        "file": "figure_7_1_graphical_abstract.png",
        "number": "7.1",
        "caption": (
            "Graphical abstract separating controlled Study A evidence from "
            "the pending independent Study B and intervention Study C evidence."
        ),
        "alt": (
            "Study A verifies controlled pipeline consistency. Study B adds "
            "independent natural-trace annotations. Study C adds counterfactual "
            "and equal-budget mitigation evidence. Study A reduces tokens but "
            "shows a task-performance degradation risk."
        ),
    },
}

EQUATION_MAP = {
    (
        r"C_i = S_i \mathbin{\|} U_i \mathbin{\|} F_i \mathbin{\|} R_i "
        r"\mathbin{\|} M_i \mathbin{\|} T_i \mathbin{\|} G_i,"
    ): "C_i = S_i || U_i || F_i || R_i || M_i || T_i || G_i,",
    (
        r"C_i = \langle m_{i,1}, m_{i,2}, \ldots, m_{i,n_i} \rangle"
    ): "C_i = <m_(i,1), m_(i,2), ..., m_(i,n_i)>",
    (
        r"S_i = \langle s_{i,1}, s_{i,2}, \ldots, s_{i,k_i} \rangle."
    ): "S_i = <s_(i,1), s_(i,2), ..., s_(i,k_i)>.",
    (
        r"J(a,b) = \frac{|W(a) \cap W(b)|}{|W(a) \cup W(b)|}."
    ): "J(a,b) = |W(a) intersect W(b)| / |W(a) union W(b)|.",
    (
        r"Q(s,q) = \frac{|W(s) \cap W(q)|}{|W(q)|}."
    ): "Q(s,q) = |W(s) intersect W(q)| / |W(q)|.",
    r"T_i = \sum_{s \in S_i} t(s).": "T_i = sum_(s in S_i) t(s).",
    (
        r"SC_{i,x} = \frac{\sum_{s \in S_i:\sigma(s)=x} t(s)}{T_i}."
    ): "SC_(i,x) = sum_(s in S_i: sigma(s)=x) t(s) / T_i.",
    r"RR_i = \frac{\sum_{s \in R_i} t(s)}{T_i}.": (
        "RR_i = sum_(s in R_i) t(s) / T_i."
    ),
    (
        r"GBR_i = \frac{\sum_{s \in G_i}t(s)}{T_i} "
        r"\quad\text{and}\quad "
        r"DBR_i = \frac{\sum_{s \in D_i}t(s)}{T_i}."
    ): (
        "GBR_i = sum_(s in G_i) t(s) / T_i, and "
        "DBR_i = sum_(s in D_i) t(s) / T_i."
    ),
    r"MBR_i = \max(RR_i, NRR_i, DBR_i).": (
        "MBR_i = max(RR_i, NRR_i, DBR_i)."
    ),
    r"CGR_i = \frac{T_i - T_{i-1}}{T_{i-1}},": (
        "CGR_i = (T_i - T_(i-1)) / T_(i-1),"
    ),
    (
        r"30\ \text{tasks} \times 6\ \text{conditions} \times "
        r"3\ \text{repetitions} \times 2\ \text{frameworks} = 1080"
    ): "30 tasks x 6 conditions x 3 repetitions x 2 frameworks = 1,080",
    (
        r"6\ \text{tasks} \times 3\ \text{conditions} \times "
        r"3\ \text{repetitions} \times 2\ \text{frameworks} = 108"
    ): "6 tasks x 3 conditions x 3 repetitions x 2 frameworks = 108",
    (
        r"\Delta T_j = T_j^{before}-T_j^{after}, \qquad "
        r"\Delta R_j = \frac{\Delta T_j}{T_j^{before}}."
    ): (
        "Delta T_j = T_j^(before) - T_j^(after),    "
        "Delta R_j = Delta T_j / T_j^(before)."
    ),
}


@dataclass
class BibEntry:
    key: str
    fields: dict[str, str]


class CitationRegistry:
    def __init__(self, entries: dict[str, BibEntry]) -> None:
        self.entries = entries
        self.order: list[str] = []

    def cite(self, keys: list[str]) -> str:
        numbers: list[int] = []
        for key in keys:
            if key not in self.entries:
                raise KeyError(f"Missing bibliography entry: {key}")
            if key not in self.order:
                self.order.append(key)
            numbers.append(self.order.index(key) + 1)
        return ", ".join(f"[{number}]" for number in numbers)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def parse_bibtex(path: Path) -> dict[str, BibEntry]:
    text = path.read_text(encoding="utf-8")
    entries: dict[str, BibEntry] = {}
    cursor = 0
    while True:
        match = re.search(r"@\w+\s*\{\s*([^,\s]+)\s*,", text[cursor:])
        if not match:
            break
        key = match.group(1)
        start = cursor + match.end()
        depth = 1
        end = start
        while end < len(text) and depth:
            if text[end] == "{":
                depth += 1
            elif text[end] == "}":
                depth -= 1
            end += 1
        body = text[start : end - 1]
        fields: dict[str, str] = {}
        field_pattern = re.compile(
            r"(\w+)\s*=\s*\{((?:[^{}]|\{[^{}]*\})*)\}\s*,?",
            flags=re.DOTALL,
        )
        for field_match in field_pattern.finditer(body):
            value = re.sub(r"\s+", " ", field_match.group(2)).strip()
            fields[field_match.group(1).lower()] = value.replace("{", "").replace(
                "}", ""
            )
        entries[key] = BibEntry(key=key, fields=fields)
        cursor = end
    return entries


def abstract_blocks(path: Path) -> list[str]:
    blocks: list[str] = []
    current: list[str] = []
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if line.startswith("# ") or line.startswith("**Status:**"):
            continue
        if line.startswith("**Keywords:**"):
            break
        if not line:
            if current:
                blocks.append(" ".join(current))
                current = []
            continue
        current.append(line)
    if current:
        blocks.append(" ".join(current))
    if len(blocks) != 5:
        raise ValueError(f"Expected five abstract paragraphs, found {len(blocks)}")
    return blocks


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_text(paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    paragraph.add_run(text)


def set_label_value(paragraph, label: str, value: str) -> None:
    clear_paragraph(paragraph)
    label_run = paragraph.add_run(label)
    label_run.bold = True
    paragraph.add_run(value)


def suppress_numbering(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    for old in p_pr.findall(qn("w:numPr")):
        p_pr.remove(old)
    num_pr = OxmlElement("w:numPr")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "0")
    num_pr.append(num_id)
    p_pr.append(num_pr)


def set_numbering(paragraph, num_id_value: int, level: int = 0) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    for old in p_pr.findall(qn("w:numPr")):
        p_pr.remove(old)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), str(num_id_value))
    num_pr.extend((ilvl, num_id))
    p_pr.append(num_pr)


def add_toc_field(paragraph) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), r'TOC \o "1-3" \h \z \u')
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "The table of contents will update when opened in Microsoft Word."
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def set_update_fields(document: Document) -> None:
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def fill_front_matter(document: Document, abstract: list[str]) -> tuple[object, object]:
    paragraphs = document.paragraphs
    if len(paragraphs) < 53:
        raise ValueError("Template does not contain the expected front matter.")

    section_breaks: list[tuple[int, object]] = []
    for index, paragraph in enumerate(paragraphs):
        p_pr = paragraph._p.find(qn("w:pPr"))
        sect_pr = p_pr.find(qn("w:sectPr")) if p_pr is not None else None
        if sect_pr is not None:
            section_breaks.append((index, copy.deepcopy(sect_pr)))
    if [index for index, _ in section_breaks[:2]] != [40, 52]:
        raise ValueError(f"Unexpected template section breaks: {section_breaks}")
    second_section_properties = section_breaks[1][1]

    set_text(paragraphs[1], TITLE)
    set_text(paragraphs[2], "")
    set_text(paragraphs[8], PROGRAMME)
    set_text(paragraphs[9], THESIS_TYPE)
    set_text(paragraphs[10], "")
    set_text(paragraphs[14], AUTHOR)
    set_text(paragraphs[17], SUPERVISOR)
    set_text(paragraphs[18], "")
    set_text(paragraphs[20], MONTH_YEAR)
    set_text(paragraphs[21], PLACE)

    set_label_value(paragraphs[28], "Title: ", TITLE)
    set_label_value(paragraphs[29], "Author: ", AUTHOR)
    set_label_value(paragraphs[30], "Month and year: ", MONTH_YEAR)
    set_label_value(
        paragraphs[31],
        "Number of pages: ",
        "To be updated before final submission",
    )
    set_label_value(paragraphs[32], "Keywords: ", KEYWORDS)
    set_text(paragraphs[33], "Abstract:")
    paragraphs[33].runs[0].bold = True

    for paragraph, block in zip(paragraphs[34:39], abstract):
        paragraph.style = "Abstract"
        set_text(paragraph, block)
    paragraphs[39].style = "Abstract"
    set_text(
        paragraphs[39],
        "Draft status: human blind review remains pending.",
    )
    paragraphs[39].runs[0].italic = True

    body = document._element.body
    for content_control in list(body.findall(qn("w:sdt"))):
        instructions = " ".join(
            node.text or ""
            for node in content_control.iter(qn("w:instrText"))
        )
        if "TOC " in instructions:
            body.remove(content_control)

    first_section_end = paragraphs[40]._p
    after_front_matter = False
    for child in list(body):
        if child is first_section_end:
            after_front_matter = True
            continue
        if after_front_matter and child.tag != qn("w:sectPr"):
            body.remove(child)
    return first_section_end, second_section_properties


def add_section_boundary(document: Document, sect_pr) -> None:
    paragraph = document.add_paragraph(style="Body Text 1")
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.append(copy.deepcopy(sect_pr))


INLINE_PATTERN = re.compile(
    r"(\[@[^\]]+\]|\*\*.+?\*\*|`.+?`|\$[^$]+\$|\*[^*]+\*)"
)


def add_inline(paragraph, text: str, citations: CitationRegistry) -> None:
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        token = match.group(0)
        if token.startswith("[@"):
            keys = [
                part.strip().lstrip("@")
                for part in token[1:-1].split(";")
                if part.strip()
            ]
            paragraph.add_run(citations.cite(keys))
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif token.startswith("$"):
            run = paragraph.add_run(normalize_equation(token[1:-1]))
            run.font.name = "Cambria Math"
            run.italic = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def add_body_paragraph(document: Document, text: str, citations: CitationRegistry):
    paragraph = document.add_paragraph(style="Body Text 1")
    add_inline(paragraph, text, citations)
    return paragraph


def normalize_equation(source: str) -> str:
    normalized = re.sub(r"\s+", " ", source).strip()
    if normalized in EQUATION_MAP:
        return EQUATION_MAP[normalized]
    replacements = {
        r"\times": "x",
        r"\Delta": "Delta",
        r"\rho": "rho",
        r"\sigma": "sigma",
        r"\ge": ">=",
        r"\in": "in",
        r"\cap": "intersect",
        r"\cup": "union",
        r"\ldots": "...",
        r"\qquad": "    ",
        r"\quad": " ",
        r"\text": "",
        "{": "",
        "}": "",
    }
    for old, new in replacements.items():
        normalized = normalized.replace(old, new)
    return normalized


def add_equation(document: Document, source: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(normalize_equation(source))
    run.font.name = "Cambria Math"
    run.font.size = Pt(11)
    run.italic = True


def markdown_table(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for index, line in enumerate(lines):
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if index == 1 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        rows.append(cells)
    return rows


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cell_margins(cell, top: int = 80, start: int = 90, bottom: int = 80, end: int = 90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_table(
    document: Document,
    rows: list[list[str]],
    caption: str,
    table_number: str,
    citations: CitationRegistry,
) -> None:
    caption_paragraph = document.add_paragraph(style="Table caption text")
    caption_paragraph.paragraph_format.keep_with_next = True
    caption_run = caption_paragraph.add_run(f"Table {table_number}. {caption}")
    caption_run.bold = True

    columns = max(len(row) for row in rows)
    normalized_rows = [row + [""] * (columns - len(row)) for row in rows]
    table = document.add_table(rows=len(normalized_rows), cols=columns)
    table.style = "Table Grid"
    table.autofit = False

    weights: list[int] = []
    for column_index in range(columns):
        maximum = max(
            len(row[column_index]) for row in normalized_rows
        )
        weights.append(max(8, min(maximum, 42)))
    total_weight = sum(weights)
    total_width = 6.55
    widths = [total_width * weight / total_weight for weight in weights]

    for row_index, values in enumerate(normalized_rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.width = Inches(widths[column_index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.style = "Table text"
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if column_index > 0 and len(value) < 24
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            clear_paragraph(paragraph)
            add_inline(paragraph, value, citations)
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E8EEF3")
                cell._tc.get_or_add_tcPr().append(shading)
    set_repeat_table_header(table.rows[0])
    document.add_paragraph(style="Body Text 1")


def font(size: int, bold: bool = False):
    candidates = (
        Path(r"C:\Windows\Fonts\arialbd.ttf")
        if bold
        else Path(r"C:\Windows\Fonts\arial.ttf")
    )
    if candidates.exists():
        return ImageFont.truetype(str(candidates), size)
    return ImageFont.load_default()


def draw_horizontal_chart(
    path: Path,
    title: str,
    labels: list[str],
    values: list[float],
    value_suffix: str = "",
) -> None:
    width, height = 1500, 210 + 90 * len(labels)
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((55, 35), title, font=font(34, True), fill="#111111")
    maximum = max(values) if max(values) > 0 else 1
    bar_left, bar_right = 470, 1320
    colors = ("#287271", "#E07A5F", "#3D5A80", "#8F6D9F", "#4C956C", "#C58C35")
    for index, (label, value) in enumerate(zip(labels, values)):
        y = 130 + index * 90
        draw.text((55, y + 6), label, font=font(24), fill="#222222")
        draw.rectangle((bar_left, y, bar_right, y + 36), fill="#EEF1F3")
        length = int((bar_right - bar_left) * value / maximum)
        draw.rectangle(
            (bar_left, y, bar_left + max(length, 2), y + 36),
            fill=colors[index % len(colors)],
        )
        value_text = f"{value:.2f}{value_suffix}"
        draw.text(
            (min(bar_left + length + 14, 1360), y + 5),
            value_text,
            font=font(22, True),
            fill="#111111",
        )
    image.save(path, dpi=(180, 180))


def draw_architecture(path: Path) -> None:
    width, height = 1700, 640
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = font(34, True)
    box_title = font(25, True)
    box_text = font(20)
    draw.text(
        (55, 30),
        "Runtime context-bloat auditing data flow",
        font=title_font,
        fill="#111111",
    )
    boxes = [
        (55, 190, 310, 390, "Agent runtime", "LangChain or\nCustom ReAct", "#DDEBF2"),
        (365, 190, 630, 390, "Instrumentation", "Capture complete\nmodel-visible input", "#E8E4F0"),
        (685, 190, 950, 390, "Provenance", "Segment, label,\ntokenize, hash", "#E5EFE5"),
        (1005, 190, 1270, 390, "Analysis", "Measure and\nlocalize bloat", "#F5E8DE"),
        (1325, 190, 1590, 390, "Mitigation", "Remove or compress\nsource-aware content", "#F0E8D8"),
    ]
    for left, top, right, bottom, heading, body, fill in boxes:
        draw.rectangle((left, top, right, bottom), fill=fill, outline="#30343B", width=3)
        heading_width = draw.textbbox((0, 0), heading, font=box_title)[2]
        draw.text(
            ((left + right - heading_width) / 2, top + 35),
            heading,
            font=box_title,
            fill="#111111",
        )
        for line_index, line in enumerate(body.splitlines()):
            line_width = draw.textbbox((0, 0), line, font=box_text)[2]
            draw.text(
                ((left + right - line_width) / 2, top + 95 + line_index * 30),
                line,
                font=box_text,
                fill="#222222",
            )
    for index in range(len(boxes) - 1):
        start = boxes[index][2] + 8
        end = boxes[index + 1][0] - 8
        y = 290
        draw.line((start, y, end, y), fill="#30343B", width=5)
        draw.polygon(
            ((end, y), (end - 18, y - 12), (end - 18, y + 12)),
            fill="#30343B",
        )
    draw.text(
        (420, 470),
        "Immutable run manifest + JSONL invocation traces + CSV/JSON evidence",
        font=font(24, True),
        fill="#30343B",
    )
    draw.text(
        (470, 520),
        "Every result remains attributable to code, configuration, data, model, and seed.",
        font=font(20),
        fill="#30343B",
    )
    image.save(path, dpi=(180, 180))


def add_figure(
    document: Document,
    image_path: Path,
    caption: str,
    figure_number: str,
    alt_text: str,
    width: float = 6.35,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    picture = paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    picture._inline.docPr.set("descr", alt_text)
    caption_paragraph = document.add_paragraph(style="Figure caption text")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    caption_paragraph.add_run(f"Figure {figure_number}. {caption}")


def parse_markdown(path: Path):
    lines = path.read_text(encoding="utf-8").splitlines()
    blocks: list[tuple[str, object]] = []
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("# "):
            blocks.append(("chapter", line[2:].strip()))
            index += 1
            continue
        if line.startswith("## "):
            blocks.append(("section", line[3:].strip()))
            index += 1
            continue
        if line == "$$":
            equation_lines: list[str] = []
            index += 1
            while index < len(lines) and lines[index].strip() != "$$":
                equation_lines.append(lines[index].strip())
                index += 1
            blocks.append(("equation", " ".join(equation_lines)))
            index += 1
            continue
        if line.startswith("```"):
            language = line[3:].strip()
            code_lines: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            blocks.append(("diagram" if language == "mermaid" else "code", code_lines))
            index += 1
            continue
        figure_match = re.fullmatch(r"\[\[FIGURE:([a-z0-9_]+)\]\]", line)
        if figure_match:
            blocks.append(("core_figure", figure_match.group(1)))
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines):
            table_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            blocks.append(("table", markdown_table(table_lines)))
            continue
        list_match = re.match(r"^(\d{1,2})\.\s+(.*)", line)
        if list_match:
            items: list[str] = []
            while index < len(lines):
                item_match = re.match(
                    r"^\s*\d{1,2}\.\s+(.*)", lines[index].strip()
                )
                if not item_match:
                    break
                items.append(item_match.group(1))
                index += 1
            blocks.append(("numbered_list", items))
            continue
        bullet_match = re.match(r"^[-*]\s+(.*)", line)
        if bullet_match:
            items = []
            while index < len(lines):
                item_match = re.match(r"^\s*[-*]\s+(.*)", lines[index].strip())
                if not item_match:
                    break
                items.append(item_match.group(1))
                index += 1
            blocks.append(("bullet_list", items))
            continue

        paragraph_lines = [line]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate:
                break
            if (
                candidate.startswith("#")
                or candidate == "$$"
                or candidate.startswith("```")
                or re.fullmatch(r"\[\[FIGURE:[a-z0-9_]+\]\]", candidate)
                or candidate.startswith("|")
                or re.match(r"^\d{1,2}\.\s+", candidate)
                or re.match(r"^[-*]\s+", candidate)
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        blocks.append(("paragraph", " ".join(paragraph_lines)))
    return blocks


def render_chapter(
    document: Document,
    path: Path,
    citations: CitationRegistry,
    assets_dir: Path,
    core_figures_dir: Path,
    counters: dict[str, int],
) -> None:
    local_table_index = 0
    for block_type, payload in parse_markdown(path):
        if block_type == "chapter":
            title = re.sub(r"^Chapter\s+\d+:\s*", "", str(payload))
            document.add_paragraph(title, style="Heading 1")
        elif block_type == "section":
            title = re.sub(r"^\d+\.\d+\s+", "", str(payload))
            document.add_paragraph(title, style="Heading 2")
        elif block_type == "paragraph":
            if re.fullmatch(r"\*Figure\s+\d+\.\d+\..*\*", str(payload)):
                continue
            add_body_paragraph(document, str(payload), citations)
        elif block_type == "equation":
            add_equation(document, str(payload))
        elif block_type == "numbered_list":
            for item in payload:
                paragraph = document.add_paragraph(style="List Paragraph")
                set_numbering(paragraph, 1)
                add_inline(paragraph, item, citations)
        elif block_type == "bullet_list":
            for item in payload:
                paragraph = document.add_paragraph(style="List Paragraph")
                set_numbering(paragraph, 2)
                add_inline(paragraph, item, citations)
        elif block_type == "code":
            for line in payload:
                paragraph = document.add_paragraph(style="Citation / Lainaus")
                run = paragraph.add_run(line)
                run.font.name = "Consolas"
                run.font.size = Pt(9)
        elif block_type == "diagram":
            add_figure(
                document,
                assets_dir / "runtime_auditing_architecture.png",
                "Runtime auditing architecture used to capture, attribute, analyze, "
                "and mitigate automatically constructed context.",
                "3.1",
                "Five-stage flow from agent runtime through instrumentation, "
                "provenance, analysis, and mitigation, with immutable trace artifacts.",
            )
        elif block_type == "core_figure":
            spec = CORE_FIGURES.get(str(payload))
            if spec is None:
                raise ValueError(f"Unknown core figure marker: {payload}")
            add_figure(
                document,
                core_figures_dir / str(spec["file"]),
                str(spec["caption"]),
                str(spec["number"]),
                str(spec["alt"]),
                width=6.45,
            )
        elif block_type == "table":
            captions = TABLE_CAPTIONS.get(path.name, [])
            caption = (
                captions[local_table_index]
                if local_table_index < len(captions)
                else f"Results reported in {path.stem}."
            )
            numbers = TABLE_NUMBERS.get(path.name, [])
            number = (
                numbers[local_table_index]
                if local_table_index < len(numbers)
                else str(counters["table"] + 1)
            )
            counters["table"] += 1
            add_table(
                document,
                payload,
                caption,
                number,
                citations,
            )
            local_table_index += 1

            if path.name == "05_results.md" and local_table_index == 2:
                add_figure(
                    document,
                    assets_dir / "tokens_by_condition.png",
                    "Mean model-visible context tokens by primary experimental condition.",
                    "5.1",
                    "Horizontal bars show mean tokens increasing from baseline and "
                    "sufficient context to source-specific and combined bloat, then "
                    "decreasing in the mitigated condition.",
                )
            if path.name == "05_results.md" and local_table_index == 3:
                add_figure(
                    document,
                    assets_dir / "source_bloat_ratio.png",
                    "Mean source-specific bloat ratio by workflow family.",
                    "5.2",
                    "Horizontal bars rank tool output first, retrieval second, and "
                    "memory third by mean source-specific bloat ratio.",
                )
            if path.name == "06_mitigation_discussion.md" and local_table_index == 1:
                add_figure(
                    document,
                    assets_dir / "mitigation_reduction.png",
                    "Mean token reduction achieved by mitigation across workflow families.",
                    "6.1",
                    "Horizontal bars show the largest reduction in tool workflows, "
                    "followed by retrieval and memory workflows.",
                )


def format_reference(number: int, entry: BibEntry) -> str:
    fields = entry.fields
    author = fields.get("author", "Unknown author").replace(" and ", ", ")
    title = fields.get("title", "Untitled")
    journal = fields.get("journal") or fields.get("howpublished", "")
    year = fields.get("year", "n.d.")
    url = fields.get("url", "")
    note = fields.get("note", "")
    parts = [f"[{number}] {author}, \"{title}.\""]
    if journal:
        parts.append(journal + ",")
    parts.append(year + ".")
    if url:
        parts.append(f"[Online]. Available: {url}.")
    if note:
        parts.append(note + ".")
    return " ".join(parts)


def add_references(
    document: Document,
    citations: CitationRegistry,
) -> None:
    heading = document.add_paragraph("References", style="Heading 1")
    suppress_numbering(heading)
    for number, key in enumerate(citations.order, start=1):
        paragraph = document.add_paragraph(style="References")
        paragraph.paragraph_format.left_indent = Inches(0.28)
        paragraph.paragraph_format.first_line_indent = Inches(-0.28)
        paragraph.add_run(format_reference(number, citations.entries[key]))


def add_unnumbered_heading(document: Document, text: str, level: int = 1):
    paragraph = document.add_paragraph(text, style=f"Heading {level}")
    suppress_numbering(paragraph)
    return paragraph


def add_appendices(document: Document, citations: CitationRegistry) -> None:
    add_unnumbered_heading(document, "Appendices", 1)

    heading = add_unnumbered_heading(
        document, "Appendix A. Context Bloat Taxonomy and Trace Schema", 2
    )
    heading.paragraph_format.page_break_before = True
    add_body_paragraph(
        document,
        "The operational taxonomy distinguishes exact duplicate, near duplicate, "
        "low-query-relevance or stale content, verbose tool output, and cross-source "
        "accumulation. Labels are attached to source-provenance segments rather than "
        "to an invocation as an indivisible string.",
        citations,
    )
    taxonomy = [
        ["Label", "Eligible sources", "Operational evidence"],
        ["exact_duplicate", "retrieval, memory, tool", "Repeated normalized hash"],
        ["near_duplicate", "retrieval, memory", "Jaccard similarity at or above 0.8"],
        [
            "low_query_relevance",
            "retrieval, memory",
            "Explicit relevance score below 0.05",
        ],
        [
            "verbose_tool_output",
            "tool",
            "Payload exceeds the compact task-specific representation",
        ],
    ]
    add_table(
        document,
        taxonomy,
        "Operational context-bloat labels.",
        "A.1",
        citations,
    )
    add_body_paragraph(
        document,
        "Each invocation trace records trace and task identifiers, framework, "
        "provider, model, configuration, workflow family, invocation index, "
        "timestamp, ordered messages and segments, character and token counts, "
        "raw and normalized hashes, source labels, injected labels, heuristic "
        "indicators, human annotations, counterfactual outcomes, aggregate "
        "metrics, usage, and risk flags.",
        citations,
    )

    heading = add_unnumbered_heading(
        document, "Appendix B. Dataset, Conditions, and Configurations", 2
    )
    heading.paragraph_format.page_break_before = True
    add_body_paragraph(
        document,
        "Dataset context_bloat_benchmark/v1 contains six calibration tasks and "
        "30 held-out primary tasks. The held-out split contains ten retrieval, ten "
        "memory, and ten tool tasks. Each task is evaluated under baseline, "
        "sufficient, exact_bloat, source_specific_bloat, combined_bloat, and "
        "mitigated conditions for both agent implementations and three repetitions.",
        citations,
    )
    config_rows = [
        ["Artifact", "Frozen value"],
        ["Model", "deepseek-v4-flash"],
        ["Frameworks", "LangChain; Custom ReAct"],
        ["Trace/project schema", "1.1.0"],
        ["Generation temperature", "0.0"],
        ["Generation seed", "20260726"],
        ["Formal bundle", "runs/studies/formal-deepseek-v1-final.zip"],
        [
            "Bundle SHA-256",
            "C6EECA7B3D588E26CC7174E0DF384754AB0BDB34E6D8B9F00788B67E46946250",
        ],
    ]
    add_table(
        document,
        config_rows,
        "Frozen formal-study configuration.",
        "B.1",
        citations,
    )

    heading = add_unnumbered_heading(
        document, "Appendix C. Additional Statistical Results", 2
    )
    heading.paragraph_format.page_break_before = True
    add_body_paragraph(
        document,
        "The frozen bundle contains 1,596 invocation traces: 1,380 primary traces "
        "and 216 cross-source stress traces. It contains 1,188 completed final task "
        "results, of which 1,080 belong to the primary matrix. The controlled "
        "detector produced 1,122 true-positive labels and no false positive or false "
        "negative labels. These values establish internal consistency for the "
        "controlled taxonomy but are not interpreted as field accuracy.",
        citations,
    )
    add_body_paragraph(
        document,
        "Task-cluster bootstrap intervals use 10,000 resamples and seed 20260726. "
        "Mitigation reduces tokens by 48.85% on average, with a 95% interval of "
        "[42.67%, 55.31%]. The paired task-success difference is -6.11 percentage "
        "points, with a 95% interval of [-12.22, -0.56]. The interval crosses the "
        "-5 percentage-point non-inferiority margin.",
        citations,
    )

    heading = add_unnumbered_heading(
        document, "Appendix D. Reproduction Instructions", 2
    )
    heading.paragraph_format.page_break_before = True
    add_body_paragraph(
        document,
        "The repository uses Python 3.11 or 3.12 and installs the package from "
        "pyproject.toml. The following commands reproduce the validation and export "
        "path from the repository root.",
        citations,
    )
    reproduction_steps = [
        "Install the project and development dependencies in an isolated environment.",
        "Run `python -m pytest -q` and require the complete suite to pass.",
        "Run `python -m context_auditor.cli validate-study --bundle "
        "runs/studies/formal-deepseek-v1-final.zip`.",
        "Verify the bundle SHA-256 and the dataset, configuration, and code hashes "
        "recorded in its study manifest.",
        "Use the exported CSV/JSON evidence files for tables and conclusions; do "
        "not recalculate formal claims from mutable ad hoc files.",
    ]
    for step in reproduction_steps:
        paragraph = document.add_paragraph(style="List Paragraph")
        set_numbering(paragraph, 1)
        add_inline(paragraph, step, citations)

    heading = add_unnumbered_heading(
        document, "Appendix E. Human Evaluation Protocol", 2
    )
    heading.paragraph_format.page_break_before = True
    add_body_paragraph(
        document,
        "The planned blind review samples 72 outputs using stratification by "
        "workflow, framework, and condition. Reviewer A assesses all 72 outputs. "
        "Reviewer B independently assesses a 36-output subset. Review files conceal "
        "framework and condition identifiers and record task success, retained "
        "evidence sufficiency, and comments. Inter-rater agreement and adjudication "
        "must be reported before the thesis is submitted. At draft v0.1, this review "
        "is pending and no human-evaluation result is claimed.",
        citations,
    )


def make_assets(assets_dir: Path) -> None:
    assets_dir.mkdir(parents=True, exist_ok=True)
    draw_architecture(assets_dir / "runtime_auditing_architecture.png")
    draw_horizontal_chart(
        assets_dir / "tokens_by_condition.png",
        "Mean context tokens by primary condition",
        [
            "Baseline",
            "Sufficient",
            "Exact bloat",
            "Source-specific bloat",
            "Combined bloat",
            "Mitigated",
        ],
        [42.70, 68.20, 76.03, 150.05, 224.73, 75.87],
        " tokens",
    )
    draw_horizontal_chart(
        assets_dir / "source_bloat_ratio.png",
        "Mean source-specific bloat ratio",
        ["Tool", "Retrieval", "Memory"],
        [50.0, 48.7, 46.8],
        "%",
    )
    draw_horizontal_chart(
        assets_dir / "mitigation_reduction.png",
        "Mean token reduction under mitigation",
        ["Retrieval QA", "Memory turns", "Multi-step tool"],
        [32.40, 29.96, 84.18],
        "%",
    )


def build(
    template: Path,
    manuscript_dir: Path,
    bibliography: Path,
    output: Path,
    assets_dir: Path,
) -> None:
    if sha256(template) != TEMPLATE_SHA256:
        raise ValueError("The retained template hash does not match artifact.md.")
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        output.unlink()
    if template.suffix.lower() == ".dotx":
        instantiate(template, output)
    elif template.suffix.lower() == ".docx":
        shutil.copy2(template, output)
    else:
        raise ValueError("The retained template must be a DOCX or DOTX file.")
    make_assets(assets_dir)
    core_figures_dir = manuscript_dir / "figures"
    build_core_figures(core_figures_dir)

    document = Document(output)
    document.core_properties.title = TITLE
    document.core_properties.author = AUTHOR
    document.core_properties.subject = (
        "Master's thesis on runtime auditing of context bloat in LLM-based agents"
    )
    document.core_properties.keywords = KEYWORDS
    document.core_properties.comments = (
        "Formal draft generated from version-controlled Markdown using "
        "Guochen_Li.docx. Human review remains pending."
    )

    _, second_section_properties = fill_front_matter(
        document, abstract_blocks(manuscript_dir / "abstract_en.md")
    )
    document.add_paragraph("Table of contents", style="TOC Heading")
    toc = document.add_paragraph()
    add_toc_field(toc)
    add_section_boundary(document, second_section_properties)

    registry = CitationRegistry(parse_bibtex(bibliography))
    counters = {"table": 0, "figure": 0}
    chapters_dir = manuscript_dir / "chapters"
    for chapter_file in CHAPTER_FILES:
        render_chapter(
            document,
            chapters_dir / chapter_file,
            registry,
            assets_dir,
            core_figures_dir,
            counters,
        )

    add_references(document, registry)
    add_appendices(document, registry)
    set_update_fields(document)
    document.save(output)

    if sha256(template) != TEMPLATE_SHA256:
        raise RuntimeError("The retained template changed during the build.")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--manuscript-dir", type=Path, required=True)
    parser.add_argument("--bibliography", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--assets-dir", type=Path, required=True)
    args = parser.parse_args()
    build(
        args.template,
        args.manuscript_dir,
        args.bibliography,
        args.output,
        args.assets_dir,
    )
    print(f"Built formal thesis draft: {args.output}")


if __name__ == "__main__":
    main()
