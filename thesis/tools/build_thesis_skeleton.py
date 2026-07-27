"""Build the version-controlled thesis skeleton from the official Word template."""

from __future__ import annotations

import argparse
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK


TITLE = (
    "Runtime Auditing of Context Bloat in LLM-Based Agents: Detection, "
    "Measurement, and Mitigation Evaluation"
)
AUTHOR = "Guochen Li"
MONTH_YEAR = "July 2026"
KEYWORDS = (
    "LLM agents; context bloat; runtime auditing; context provenance; "
    "retrieval-augmented generation; agent memory; tool use"
)
CCS = (
    "ACM Computing Classification System (CCS)\n"
    "Computing methodologies -> Artificial intelligence -> "
    "Natural language processing"
)

CHAPTER_FILES = (
    "01_introduction.md",
    "02_background_related_work.md",
    "03_model_and_framework.md",
    "04_methodology.md",
    "05_results.md",
    "06_mitigation_discussion.md",
    "07_conclusion.md",
)

APPENDICES = (
    (
        "Appendix A. Context Bloat Taxonomy and Trace Schema",
        "The final appendix will contain the frozen taxonomy, source types, "
        "trace schema, and field definitions required to interpret the artifact.",
    ),
    (
        "Appendix B. Dataset, Conditions, and Configurations",
        "The final appendix will document the dataset version, task allocation, "
        "six experimental conditions, model parameters, and run configurations.",
    ),
    (
        "Appendix C. Additional Statistical Results",
        "The final appendix will contain extended tables, confidence intervals, "
        "label-level results, and stress-cohort analyses omitted from the main text.",
    ),
    (
        "Appendix D. Reproduction Instructions",
        "The final appendix will provide environment, validation, execution, "
        "bundle-export, and checksum-verification instructions.",
    ),
    (
        "Appendix E. Human Evaluation Protocol",
        "The final appendix will describe sampling, blinding, rating fields, "
        "reliability analysis, and the status of the human review.",
    ),
)


def _clear_paragraph(paragraph) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)


def _replace_paragraph_text(paragraph, text: str) -> None:
    _clear_paragraph(paragraph)
    paragraph.add_run(text)


def _replace_exact_paragraph(document: Document, old: str, new: str) -> None:
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text == old]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected exactly one paragraph containing {old!r}, found {len(matches)}."
        )
    _replace_paragraph_text(matches[0], new)


def _set_cell_value(cell, value: str) -> None:
    if len(cell.paragraphs) < 2:
        paragraph = cell.add_paragraph()
    else:
        paragraph = cell.paragraphs[1]
    paragraph.style = "Abstract table paragraph"
    _replace_paragraph_text(paragraph, value)
    for extra in list(cell.paragraphs[2:]):
        extra._element.getparent().remove(extra._element)


def _read_abstract(path: Path) -> list[str]:
    blocks: list[str] = []
    current: list[str] = []

    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if line.startswith("# ") or line.startswith("**Status:**"):
            continue
        if line.startswith("**Keywords:**"):
            break
        if not line:
            if current:
                blocks.append(" ".join(current))
                current = []
            continue
        current.append(line)

    if current:
        blocks.append(" ".join(current))
    if len(blocks) != 5:
        raise RuntimeError(f"Expected five abstract paragraphs, found {len(blocks)}.")
    return blocks


def _fill_abstract_table(document: Document, abstract_blocks: list[str]) -> None:
    table = document.tables[0]
    _set_cell_value(table.rows[0].cells[0], "Faculty of Science")
    _set_cell_value(table.rows[0].cells[2], "Master's Programme in Computer Science")
    _set_cell_value(table.rows[1].cells[0], AUTHOR)
    _set_cell_value(table.rows[2].cells[0], TITLE)
    _set_cell_value(table.rows[3].cells[0], "[Supervisor name to be confirmed]")
    _set_cell_value(table.rows[4].cells[0], "Master's thesis")
    _set_cell_value(table.rows[4].cells[1], MONTH_YEAR)
    _set_cell_value(
        table.rows[4].cells[3], "To be completed at final submission"
    )

    abstract_cell = table.rows[5].cells[0]
    for paragraph in list(abstract_cell.paragraphs[1:]):
        paragraph._element.getparent().remove(paragraph._element)
    for block in abstract_blocks:
        paragraph = abstract_cell.add_paragraph(style="Abstract text")
        paragraph.add_run(block)
    ccs_paragraph = abstract_cell.add_paragraph(style="Abstract text")
    ccs_lines = CCS.splitlines()
    heading = ccs_paragraph.add_run(ccs_lines[0])
    heading.bold = True
    ccs_paragraph.add_run("\n" + ccs_lines[1])

    _set_cell_value(table.rows[6].cells[0], KEYWORDS)
    _set_cell_value(table.rows[7].cells[0], "Helsinki University Library")
    _set_cell_value(
        table.rows[8].cells[0],
        "Draft v0.1; human blind review pending; study track to be confirmed",
    )


def _remove_content_after_contents(document: Document) -> None:
    contents = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip() == "Contents"
    ]
    if len(contents) != 1:
        raise RuntimeError(f"Expected one Contents heading, found {len(contents)}.")

    body = document._element.body
    found = False
    for child in list(body):
        if child is contents[0]._p:
            found = True
            continue
        if found and child.tag != document._element.body.sectPr.tag:
            body.remove(child)


def _flush_block(blocks: list[str], current: list[str]) -> None:
    if current:
        blocks.append(" ".join(current))
        current.clear()


def _parse_chapter(path: Path) -> tuple[str, list[tuple[str, list[str]]]]:
    chapter_title = ""
    sections: list[tuple[str, list[str]]] = []
    section_title = ""
    blocks: list[str] = []
    current: list[str] = []

    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if line.startswith("# "):
            chapter_title = re.sub(r"^Chapter\s+\d+:\s*", "", line[2:])
            continue
        if line.startswith("## "):
            _flush_block(blocks, current)
            if section_title:
                sections.append((section_title, blocks))
            section_title = re.sub(r"^\d+\.\d+\s+", "", line[3:])
            blocks = []
            continue
        if line.startswith(">"):
            continue
        if not line:
            _flush_block(blocks, current)
            continue
        current.append(line)

    _flush_block(blocks, current)
    if section_title:
        sections.append((section_title, blocks))
    if not chapter_title or not sections:
        raise RuntimeError(f"Could not parse chapter skeleton: {path}")
    return chapter_title, sections


def _add_page_break(document: Document) -> None:
    paragraph = document.add_paragraph(style="Text")
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def _add_drafting_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Text")
    label = paragraph.add_run("Drafting brief. ")
    label.bold = True
    content = paragraph.add_run(text)
    content.italic = True


def _add_manuscript_body(document: Document, chapters_dir: Path) -> None:
    toc_placeholder = document.add_paragraph("[[TOC]]", style="Text")
    toc_placeholder.paragraph_format.keep_with_next = False

    for chapter_file in CHAPTER_FILES:
        chapter_title, sections = _parse_chapter(chapters_dir / chapter_file)
        _add_page_break(document)
        document.add_paragraph(chapter_title, style="Heading 1")
        for section_title, blocks in sections:
            document.add_paragraph(section_title, style="Heading 2")
            for block in blocks:
                _add_drafting_paragraph(document, block)

    _add_page_break(document)
    document.add_paragraph("References", style="Heading no number")
    document.add_paragraph(
        "The reference list will be generated from the version-controlled "
        "bibliography and checked against every in-text citation.",
        style="Bibtext",
    )

    for appendix_title, description in APPENDICES:
        _add_page_break(document)
        document.add_paragraph(appendix_title, style="Heading no number")
        _add_drafting_paragraph(document, description)


def build(template: Path, manuscript_dir: Path, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    if template.resolve() == output.resolve():
        raise RuntimeError("Output must not overwrite the retained template.")
    shutil.copy2(template, output)

    document = Document(output)
    document.core_properties.title = TITLE
    document.core_properties.author = AUTHOR
    document.core_properties.subject = "Master's thesis skeleton"
    document.core_properties.comments = (
        "Generated from version-controlled Markdown sources and the retained "
        "University of Helsinki Computer Science thesis template."
    )

    _replace_exact_paragraph(document, "Title", TITLE)
    _replace_exact_paragraph(document, "Firstname Lastname ", AUTHOR)
    _replace_exact_paragraph(document, "September 5, 2021", MONTH_YEAR)
    _fill_abstract_table(
        document, _read_abstract(manuscript_dir / "abstract_en.md")
    )
    _remove_content_after_contents(document)
    _add_manuscript_body(document, manuscript_dir / "chapters")
    document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--manuscript-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.template, args.manuscript_dir, args.output)
    print(f"Built thesis skeleton: {args.output}")


if __name__ == "__main__":
    main()
