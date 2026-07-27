"""Small, resumable Excel workbooks for blinded human annotation."""

from __future__ import annotations

import csv
from copy import copy
from pathlib import Path
from typing import Iterable


def write_annotation_workbook(
    path: str | Path,
    *,
    fields: tuple[str, ...],
    rows: Iterable[dict[str, object]],
    block_id: str,
    decision_values: tuple[str, ...],
    confidence_field: str = "confidence_1_to_5",
) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError as error:
        raise RuntimeError(
            "Excel annotation packages require openpyxl==3.1.5"
        ) from error
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    review = workbook.active
    review.title = "Review"
    review.append(list(fields))
    materialized = list(rows)
    for row in materialized:
        review.append([row.get(field, "") for field in fields])
    review.freeze_panes = "A2"
    review.auto_filter.ref = review.dimensions
    for cell in review[1]:
        font = copy(cell.font)
        font.bold = True
        cell.font = font
    widths = {
        "task_prompt": 55,
        "segment_text": 80,
        "model_output": 80,
        "expected_answer": 35,
        "notes": 35,
        "reasons": 28,
    }
    for index, field in enumerate(fields, start=1):
        review.column_dimensions[review.cell(1, index).column_letter].width = (
            widths.get(field, max(14, min(28, len(field) + 4)))
        )
    if materialized and "decision" in fields:
        column = fields.index("decision") + 1
        validation = DataValidation(
            type="list",
            formula1='"' + ",".join(decision_values) + '"',
            allow_blank=False,
        )
        review.add_data_validation(validation)
        validation.add(
            f"{review.cell(2, column).coordinate}:"
            f"{review.cell(len(materialized) + 1, column).coordinate}"
        )
    if materialized and confidence_field in fields:
        column = fields.index(confidence_field) + 1
        confidence = DataValidation(
            type="whole",
            operator="between",
            formula1="1",
            formula2="5",
            allow_blank=False,
        )
        review.add_data_validation(confidence)
        confidence.add(
            f"{review.cell(2, column).coordinate}:"
            f"{review.cell(len(materialized) + 1, column).coordinate}"
        )
    session = workbook.create_sheet("Session")
    session.append(["block_id", block_id])
    session.append(["started_at_utc", ""])
    session.append(["completed_at_utc", ""])
    session.append(["reviewer_notes", ""])
    session.column_dimensions["A"].width = 24
    session.column_dimensions["B"].width = 42
    workbook.save(target)
    return target


def read_annotation_workbooks(
    path: str | Path,
) -> list[dict[str, str]]:
    source = Path(path)
    if source.is_dir():
        files = sorted(source.glob("block_*.xlsx"))
        if not files:
            raise ValueError(f"No block workbooks found in {source}")
        rows = []
        for file in files:
            rows.extend(read_annotation_workbook(file))
        return rows
    if source.suffix.lower() == ".xlsx":
        return read_annotation_workbook(source)
    with source.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def read_annotation_workbook(path: Path) -> list[dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError(
            "Excel annotation packages require openpyxl==3.1.5"
        ) from error
    workbook = load_workbook(path, data_only=True)
    if "Review" not in workbook.sheetnames or "Session" not in workbook.sheetnames:
        raise ValueError(f"Invalid annotation workbook: {path}")
    session = workbook["Session"]
    metadata = {
        str(session.cell(row, 1).value or ""): str(
            session.cell(row, 2).value or ""
        ).strip()
        for row in range(1, session.max_row + 1)
    }
    block_id = metadata.get("block_id", "")
    started = metadata.get("started_at_utc", "")
    completed = metadata.get("completed_at_utc", "")
    if not block_id or not started or not completed:
        raise ValueError(
            f"Block session metadata is incomplete in {path.name}"
        )
    review = workbook["Review"]
    fields = [str(cell.value or "") for cell in review[1]]
    rows: list[dict[str, str]] = []
    for values in review.iter_rows(min_row=2, values_only=True):
        if not any(value not in (None, "") for value in values):
            continue
        row = {
            field: str(value if value is not None else "")
            for field, value in zip(fields, values)
        }
        row["block_id"] = block_id
        row["block_started_at"] = started
        row["block_completed_at"] = completed
        rows.append(row)
    return rows
